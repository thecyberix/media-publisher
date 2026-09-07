from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from catalog_parser.canva_selection import (
    extract_canva_links_from_docx,
    extract_canva_links_from_google_document,
    select_canva_url,
)
from catalog_parser.drive_video_size import (
    video_size_from_drive_file_metadata,
    video_size_from_pkg_folder,
)


def _add_field_code_hyperlink(paragraph, url: str, display: str = "Canva") -> None:
    """Insert a classic Word HYPERLINK field (w:instrText), not w:hyperlink."""
    begin = OxmlElement("w:r")
    begin_char = OxmlElement("w:fldChar")
    begin_char.set(qn("w:fldCharType"), "begin")
    begin.append(begin_char)
    paragraph._p.append(begin)

    instr = OxmlElement("w:r")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = f' HYPERLINK "{url}" '
    instr.append(instr_text)
    paragraph._p.append(instr)

    sep = OxmlElement("w:r")
    sep_char = OxmlElement("w:fldChar")
    sep_char.set(qn("w:fldCharType"), "separate")
    sep.append(sep_char)
    paragraph._p.append(sep)

    text_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = display
    text_run.append(text)
    paragraph._p.append(text_run)

    end = OxmlElement("w:r")
    end_char = OxmlElement("w:fldChar")
    end_char.set(qn("w:fldCharType"), "end")
    end.append(end_char)
    paragraph._p.append(end)


class DocxCanvaExtractionTests(unittest.TestCase):
    def test_extracts_field_code_hyperlink(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Thumbnail")
        canva_url = "https://www.canva.com/design/DAFieldCode123/view?utm=1"
        _add_field_code_hyperlink(paragraph, canva_url)
        urls, _below = extract_canva_links_from_docx(document)
        self.assertEqual(urls, ["https://www.canva.com/design/DAFieldCode123"])

    def test_extracts_canva_link_shortlink_from_field_code(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _add_field_code_hyperlink(paragraph, "https://canva.link/rbgbets4hffvol0")
        with patch(
            "media_publisher.sources.canva.resolve_canva_url",
            return_value="https://www.canva.com/design/DAHKegUvggY/view",
        ):
            urls, _below = extract_canva_links_from_docx(document)
        self.assertEqual(urls, ["https://www.canva.com/design/DAHKegUvggY"])

    def test_extracts_field_code_hyperlink_in_table_cell(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "THUMBNAIL - YT"
        cell_paragraph = table.cell(1, 0).paragraphs[0]
        canva_url = "https://www.canva.com/design/DATableField99/edit"
        _add_field_code_hyperlink(cell_paragraph, canva_url, display="Open design")
        urls, _below = extract_canva_links_from_docx(document)
        self.assertEqual(urls, ["https://www.canva.com/design/DATableField99"])

    def test_extracts_modern_w_hyperlink(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        canva_url = "https://www.canva.com/design/DAModern456/view"
        rid = paragraph.part.relate_to(canva_url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "design"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        urls, _below = extract_canva_links_from_docx(document)
        self.assertEqual(urls, ["https://www.canva.com/design/DAModern456"])

    def test_roundtrip_save_keeps_field_code(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "pkg.docx"
            document = Document()
            paragraph = document.add_paragraph()
            _add_field_code_hyperlink(
                paragraph,
                "https://www.canva.com/design/DASaved789/view",
            )
            document.save(path)
            loaded = Document(str(path))
            urls, _below = extract_canva_links_from_docx(loaded)
        self.assertEqual(urls, ["https://www.canva.com/design/DASaved789"])


class GoogleDocCanvaExtractionTests(unittest.TestCase):
    def test_extracts_hyperlink_inside_table_cell(self) -> None:
        document = {
            "body": {
                "content": [
                    {
                        "table": {
                            "tableRows": [
                                {
                                    "tableCells": [
                                        {
                                            "content": [
                                                {
                                                    "paragraph": {
                                                        "elements": [
                                                            {
                                                                "textRun": {
                                                                    "content": "Open ",
                                                                    "textStyle": {},
                                                                }
                                                            },
                                                            {
                                                                "textRun": {
                                                                    "content": "design",
                                                                    "textStyle": {
                                                                        "link": {
                                                                            "url": (
                                                                                "https://www.canva.com/"
                                                                                "design/DAGoogleTbl/view"
                                                                            )
                                                                        }
                                                                    },
                                                                }
                                                            },
                                                        ]
                                                    }
                                                }
                                            ]
                                        }
                                    ]
                                }
                            ]
                        }
                    }
                ]
            }
        }
        urls, _below = extract_canva_links_from_google_document(document)
        self.assertEqual(urls, ["https://www.canva.com/design/DAGoogleTbl"])


class CanvaSelectionTests(unittest.TestCase):
    def test_select_single_url_without_probe(self) -> None:
        url = "https://www.canva.com/design/ABC123/view"
        self.assertEqual(
            select_canva_url(
                [url],
                target_size=(1920, 1080),
            ),
            "https://www.canva.com/design/ABC123",
        )

    def test_select_prefers_aspect_matching_design(self) -> None:
        landscape = "https://www.canva.com/design/LAND/view"
        portrait = "https://www.canva.com/design/PORT/view"
        with patch(
            "catalog_parser.canva_selection.probe_canva_design_dimensions",
            side_effect=lambda url: {
                "https://www.canva.com/design/LAND": (1920, 1080),
                "https://www.canva.com/design/PORT": (1080, 1920),
            }[url],
        ):
            selected = select_canva_url(
                [portrait, landscape],
                target_size=(1920, 1080),
            )
        self.assertEqual(selected, "https://www.canva.com/design/LAND")

    def test_select_prefers_drive_target_size_over_original_video_url(self) -> None:
        landscape = "https://www.canva.com/design/LAND/view"
        portrait = "https://www.canva.com/design/PORT/view"
        with patch(
            "catalog_parser.canva_selection.video_size_from_source_url",
            return_value=(1080, 1920),
        ):
            with patch(
                "catalog_parser.canva_selection.probe_canva_design_dimensions",
                side_effect=lambda url: {
                    "https://www.canva.com/design/LAND": (1920, 1080),
                    "https://www.canva.com/design/PORT": (1080, 1920),
                }[url],
            ):
                selected = select_canva_url(
                    [portrait, landscape],
                    target_size=(1920, 1080),
                    original_video_url="https://instagram.com/reel/short",
                )
        self.assertEqual(selected, "https://www.canva.com/design/LAND")


class DriveVideoSizeTests(unittest.TestCase):
    def test_video_size_from_drive_file_metadata(self) -> None:
        drive_service = MagicMock()
        drive_service.files().get().execute.return_value = {
            "videoMediaMetadata": {"width": 1920, "height": 1080},
        }
        self.assertEqual(
            video_size_from_drive_file_metadata(drive_service, "file-1"),
            (1920, 1080),
        )

    def test_video_size_from_drive_file_prefers_ffprobe_over_metadata(self) -> None:
        from catalog_parser.drive_video_size import video_size_from_drive_file

        drive_service = MagicMock()
        drive_service.files().get().execute.return_value = {
            "videoMediaMetadata": {"width": 1920, "height": 1080},
        }
        with patch(
            "catalog_parser.drive_video_size.download_drive_file",
            side_effect=lambda _drive, _file_id, destination: destination.write_bytes(
                b"fake"
            )
            or destination,
        ), patch(
            "media_publisher.video_duration.probe_local_video_size",
            return_value=(1080, 1920),
        ):
            self.assertEqual(
                video_size_from_drive_file(
                    drive_service,
                    "file-1",
                    file_name="phone.mp4",
                ),
                (1080, 1920),
            )

    def test_video_size_from_pkg_folder(self) -> None:
        drive_service = MagicMock()
        with patch(
            "catalog_parser.drive_video_size.find_video_and_audio_subfolder",
            return_value=MagicMock(video=MagicMock(id="video-1", name="All Video.mp4")),
        ):
            with patch(
                "catalog_parser.drive_video_size.video_size_from_drive_file",
                return_value=(1080, 1920),
            ):
                self.assertEqual(
                    video_size_from_pkg_folder(drive_service, "folder-1"),
                    (1080, 1920),
                )


if __name__ == "__main__":
    unittest.main()
